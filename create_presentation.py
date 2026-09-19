import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette matching user's template
    BG_CREAM = RGBColor(250, 248, 242)       # Warm ivory
    DARK_GREEN = RGBColor(13, 107, 72)       # Forest emerald
    LIGHT_GREEN = RGBColor(228, 243, 233)     # Sage card fill
    BORDER_GREEN = RGBColor(52, 168, 83)     # Green border
    DARK_SLATE = RGBColor(31, 41, 55)        # Charcoal text
    TEXT_MUTED = RGBColor(75, 85, 99)        # Medium gray
    WHITE = RGBColor(255, 255, 255)
    GOLD = RGBColor(234, 179, 8)
    ROSE = RGBColor(225, 29, 72)

    def apply_background(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_CREAM

    def add_card(slide, left, top, width, height, bg_color=LIGHT_GREEN, border_color=BORDER_GREEN):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # ==========================================
    # SLIDE 1: Title
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    # Eyebrow
    tx1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(0.6))
    tf1 = tx1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "CORRELATION, VOLATILITY & AI-POWERED RISK ANALYSIS"
    p1.font.name = "Arial"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = DARK_GREEN

    # Title
    tx_title = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(2.2))
    tf_title = tx_title.text_frame
    p_t = tf_title.paragraphs[0]
    p_t.text = "CAPITALX: PORTFOLIO\nOPTIMIZATION & AI TERMINAL"
    p_t.font.name = "Arial Black"
    p_t.font.size = Pt(40)
    p_t.font.bold = True
    p_t.font.color.rgb = DARK_GREEN

    # Subtitle
    tx_sub = s1.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(11.3), Inches(1.5))
    tf_sub = tx_sub.text_frame
    p_s = tf_sub.paragraphs[0]
    p_s.text = "Mathematical Mean-Variance Optimization, Spectral Risk Analysis,\nand Plain-English AI Fiduciary Diagnostics for Smarter Investing."
    p_s.font.name = "Arial"
    p_s.font.size = Pt(18)
    p_s.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 2: Problems & Solution
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)

    # Heading Left
    th2 = s2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(5.0), Inches(1.2))
    p = th2.text_frame.paragraphs[0]
    p.text = "PROBLEMS\nEXISTING"
    p.font.name = "Arial Black"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    # Heading Right
    th2_r = s2.shapes.add_textbox(Inches(6.8), Inches(0.8), Inches(5.5), Inches(1.2))
    p = th2_r.text_frame.paragraphs[0]
    p.text = "OUR\nSOLUTION"
    p.font.name = "Arial Black"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    # Problem Card
    c_prob = add_card(s2, Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.5), WHITE, BORDER_GREEN)
    tf_p = c_prob.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = Inches(0.3)
    tf_p.margin_right = Inches(0.3)
    tf_p.margin_top = Inches(0.3)

    probs = [
        ("01 — Hidden Correlation", "Investors hold multiple stocks (e.g. AAPL, MSFT, NVDA) that move in lockstep without realizing it."),
        ("02 — False Diversification", "Holding 5 tech stocks only provides ~1.3 Effective Bets, leaving capital vulnerable to sector crashes."),
        ("03 — Inefficient Capital Allocation", "Investors allocate money based on intuition or equal-weighting rather than risk-return efficiency."),
        ("04 — Lack of Plain-English Translation", "Existing tools dump raw covariance numbers, kurtosis, and Greek formulas without actionable trade advice.")
    ]
    for i, (title, desc) in enumerate(probs):
        p_t = tf_p.add_paragraph() if i > 0 else tf_p.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.space_before = Pt(8) if i > 0 else Pt(0)

        p_d = tf_p.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Arial"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = DARK_SLATE

    # Solution Card
    c_sol = add_card(s2, Inches(6.8), Inches(2.2), Inches(5.5), Inches(4.5), LIGHT_GREEN, BORDER_GREEN)
    tf_s = c_sol.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = Inches(0.35)
    tf_s.margin_right = Inches(0.35)
    tf_s.margin_top = Inches(0.35)

    p_st = tf_s.paragraphs[0]
    p_st.text = "CapitalX Portfolio Intelligence & Risk Terminal"
    p_st.font.name = "Arial"
    p_st.font.size = Pt(14)
    p_st.font.bold = True
    p_st.font.color.rgb = DARK_GREEN
    p_st.space_after = Pt(10)

    sols = [
        "• Quantitative Optimization Engine: Solves for the Maximum Sharpe Ratio (Tangency) and Minimum Volatility portfolios using convex quadratic programming.",
        "• Spectral Diversification (N_eff): Uses Shannon entropy of correlation eigenvalues to detect false diversification.",
        "• AI CIO & Plain-English Diagnostics: Automatically generates a fiduciary executive memo and a humorous 'Roast My Portfolio' breakdown.",
        "• 1-Click Crisis Simulator & What-If Sandbox: Tests allocations against historical market crashes (COVID, 2022 rate hike) with 1-click hedge injections.",
        "• Actionable Trade Ticket: Generates copyable BUY/SELL/HOLD order execution instructions."
    ]
    for s_text in sols:
        p = tf_s.add_paragraph()
        p.text = s_text
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.color.rgb = DARK_SLATE
        p.space_after = Pt(6)

    # ==========================================
    # SLIDE 3: Technology Stack & Target Users
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)

    th3 = s3.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(6.0), Inches(1.0))
    p = th3.text_frame.paragraphs[0]
    p.text = "TECHNOLOGY\nSTACK"
    p.font.name = "Arial Black"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    th3_u = s3.shapes.add_textbox(Inches(8.0), Inches(0.6), Inches(4.5), Inches(1.0))
    p = th3_u.text_frame.paragraphs[0]
    p.text = "TARGET\nUSERS"
    p.font.name = "Arial Black"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    # Tech Stack Grid (6 mini cards)
    tech_items = [
        ("DATA SOURCE", "yfinance", "• Historical OHLCV data\n• Automated daily log returns"),
        ("COMPUTATION", "Python, NumPy & Pandas", "• Log returns & Ledoit-Wolf cov\n• Eigenvalue decomposition"),
        ("OPTIMIZATION", "SciPy & CVXPY", "• Quadratic programming (SLSQP/OSQP)\n• Charnes-Cooper Tangency solver"),
        ("BACKEND ENGINE", "FastAPI (Async)", "• High-performance REST endpoints\n• Robust 3-tier solver fallbacks"),
        ("AI & RISK ANALYTICS", "Cornish-Fisher & Spectral Engine", "• VaR 95%, CVaR, Sortino ratio\n• Plain-English AI memo generator"),
        ("FRONTEND & EXPORT", "HTML5, Vanilla JS & Chart.js", "• Sub-50ms reactive charts\n• 1-Click PDF Tear Sheet print engine")
    ]

    grid_x = [Inches(1.0), Inches(4.4)]
    grid_y = [Inches(1.8), Inches(3.6), Inches(5.4)]

    for idx, (category, name, bullets) in enumerate(tech_items):
        col = idx % 2
        row = idx // 2
        c = add_card(s3, grid_x[col], grid_y[row], Inches(3.2), Inches(1.6), WHITE, BORDER_GREEN)
        tf = c.text_frame
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = category
        p.font.name = "Arial"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        
        p2 = tf.add_paragraph()
        p2.text = name
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = DARK_SLATE
        
        p3 = tf.add_paragraph()
        p3.text = bullets
        p3.font.name = "Arial"
        p3.font.size = Pt(8)
        p3.font.color.rgb = TEXT_MUTED

    # Target Users Card Right
    c_users = add_card(s3, Inches(8.0), Inches(1.8), Inches(4.3), Inches(5.2), LIGHT_GREEN, BORDER_GREEN)
    tf_u = c_users.text_frame
    tf_u.margin_left = Inches(0.25)
    tf_u.margin_top = Inches(0.25)

    users = [
        ("RETAIL INVESTORS", "Analyze existing stock holdings, uncover hidden correlations, and generate exact rebalance orders."),
        ("BEGINNER INVESTORS", "Understand risk, return, and diversification visually with an intuitive Health Score and Roast Mode."),
        ("STUDENTS & RESEARCHERS", "Learn real-world applications of Modern Portfolio Theory, Ledoit-Wolf shrinkage, and spectral factor analysis."),
        ("WEALTH ADVISORS (RIAs)", "Generate client-ready 1-page institutional tear sheets and stress-test portfolios against historical crises.")
    ]
    for i, (u_title, u_desc) in enumerate(users):
        p_t = tf_u.add_paragraph() if i > 0 else tf_u.paragraphs[0]
        p_t.text = u_title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(11.5)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.space_before = Pt(8) if i > 0 else Pt(0)

        p_d = tf_u.add_paragraph()
        p_d.text = u_desc
        p_d.font.name = "Arial"
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 4: Workflow and Methodology
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)

    tx4_h = s4.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.2))
    p = tx4_h.text_frame.paragraphs[0]
    p.text = "WORKFLOW AND METHODOLOGY\nHOW THE TOOL WORKS"
    p.font.name = "Arial Black"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    # 5 Horizontal Workflow Steps
    steps = [
        ("01 USER INPUT", "Enter tickers (stocks, ETFs, gold, crypto), initial weights, and lookback window (1-5 yrs)."),
        ("02 DATA INGESTION", "Fetch historical closing prices via yfinance & calculate continuous daily log returns."),
        ("03 RISK ANALYSIS", "Compute Ledoit-Wolf covariance, Pearson correlation, and Effective Bets (Shannon entropy)."),
        ("04 OPTIMIZATION", "Run convex quadratic solvers (CVXPY/SciPy) for Tangency (Max Sharpe) & GMV portfolios."),
        ("05 DIAGNOSTICS", "Generate Health Score (0-100), AI CIO Memo, Crisis Stress Replay, and Trade Tickets.")
    ]

    card_w = Inches(2.15)
    card_gap = Inches(0.14)
    start_x = Inches(1.0)

    for i, (s_title, s_desc) in enumerate(steps):
        x = start_x + i * (card_w + card_gap)
        c = add_card(s4, x, Inches(2.4), card_w, Inches(4.3), WHITE, BORDER_GREEN)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.3)

        p_num = tf.paragraphs[0]
        p_num.text = f"STEP 0{i+1}"
        p_num.font.name = "Arial Black"
        p_num.font.size = Pt(14)
        p_num.font.color.rgb = GOLD
        p_num.space_after = Pt(8)

        p_name = tf.add_paragraph()
        p_name.text = s_title.split(' ', 1)[1]
        p_name.font.name = "Arial"
        p_name.font.size = Pt(12)
        p_name.font.bold = True
        p_name.font.color.rgb = DARK_GREEN
        p_name.space_after = Pt(12)

        p_desc = tf.add_paragraph()
        p_desc.text = s_desc
        p_desc.font.name = "Arial"
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 5: Core Quantitative Outputs
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5)

    tx5_h = s5.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.2))
    p = tx5_h.text_frame.paragraphs[0]
    p.text = "CORE QUANTITATIVE ENGINE\nWHAT THE USER GETS"
    p.font.name = "Arial Black"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    core_features = [
        ("CORRELATION HEATMAP", "Interactive color-coded grid revealing pairwise Pearson correlations and identifying hidden asset collinearity."),
        ("ASSET VOLATILITY ANALYSIS", "Compares individual annualized asset volatilities against current and optimal portfolio volatility."),
        ("CURRENT VS OPTIMIZED", "Side-by-side metric comparison: Expected Return, Volatility, Sharpe Ratio, and Sortino Ratio."),
        ("EFFICIENT FRONTIER PLOT", "Interactive Markowitz scatter plot visualizing the entire hyperbola of efficient portfolios with tangent ray."),
        ("ROLLING STRESS VOLATILITY", "63-day rolling volatility time-series tracking how portfolio risk behaved during historical turbulence."),
        ("ACTIONABLE REBALANCE TABLE", "Exact percentage adjustments (Delta w) with clear BUY, SELL, or HOLD execution tags.")
    ]

    c_w = Inches(5.4)
    c_h = Inches(1.5)
    xs = [Inches(1.0), Inches(6.8)]
    ys = [Inches(2.0), Inches(3.7), Inches(5.4)]

    for idx, (f_title, f_desc) in enumerate(core_features):
        col = idx % 2
        row = idx // 2
        c = add_card(s5, xs[col], ys[row], c_w, c_h, WHITE, BORDER_GREEN)
        tf = c.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)
        p1 = tf.paragraphs[0]
        p1.text = f_title
        p1.font.name = "Arial"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = DARK_GREEN
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.name = "Arial"
        p2.font.size = Pt(10)
        p2.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 6: NEW AI INTELLIGENCE & CRISIS SUITE (The Gamechanger!)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)

    tx6_h = s6.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.2))
    p = tx6_h.text_frame.paragraphs[0]
    p.text = "NEW HACKATHON INNOVATIONS\nAI INTELLIGENCE & CRISIS SUITE"
    p.font.name = "Arial Black"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    ai_features = [
        ("PORTFOLIO HEALTH SCORE (0-100)", "Multi-pillar rating (Grade A+ to F) evaluating Diversification (N_eff), Sharpe Efficiency, Volatility Discipline, and Tail Resilience."),
        ("DUAL-VOICE AI COMMENTARY", "Toggle between an Institutional Fiduciary Memo (BlackRock-style risk thesis) and a hilarious 'Roast My Portfolio' mode exposing retail habits."),
        ("1-CLICK MACRO CRISIS SIMULATOR", "Replays the portfolio against 2020 COVID, 2022 Fed Rate Shock, 2008 Lehman GFC, and 1970s Stagflation with exact drawdown cushions."),
        ("INTERACTIVE 'WHAT-IF' SANDBOX", "Quick-injection testing of Gold (+GLD), Treasuries (+TLT), Bitcoin (+BTC), and Cash with 1-click portfolio table injection."),
        ("1-CLICK BROKER TRADE TICKET", "Generates formatted, copyable BUY/SELL orders ready to paste directly into any brokerage terminal."),
        ("INSTITUTIONAL TEAR SHEET EXPORT", "1-Click PDF print export that transforms the dashboard into a clean, 1-page institutional investment deck.")
    ]

    for idx, (f_title, f_desc) in enumerate(ai_features):
        col = idx % 2
        row = idx // 2
        c = add_card(s6, xs[col], ys[row], c_w, c_h, LIGHT_GREEN, BORDER_GREEN)
        tf = c.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)
        p1 = tf.paragraphs[0]
        p1.text = "★ " + f_title
        p1.font.name = "Arial"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = DARK_GREEN
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.name = "Arial"
        p2.font.size = Pt(10)
        p2.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 7: Future Scope & Project Value
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7)

    tx7_l = s7.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(5.5), Inches(1.0))
    p = tx7_l.text_frame.paragraphs[0]
    p.text = "FUTURE\nENHANCEMENTS"
    p.font.name = "Arial Black"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    tx7_r = s7.shapes.add_textbox(Inches(6.8), Inches(0.8), Inches(5.5), Inches(1.0))
    p = tx7_r.text_frame.paragraphs[0]
    p.text = "PROJECT\nVALUE & IMPACT"
    p.font.name = "Arial Black"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    # Left: Future Scope (Updated since we already built ETFs/Bonds/Gold/Risk measures!)
    c_fut = add_card(s7, Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.7), WHITE, BORDER_GREEN)
    tf_f = c_fut.text_frame
    tf_f.margin_left = Inches(0.3)
    tf_f.margin_top = Inches(0.3)

    future_items = [
        ("• Direct Broker Execution APIs", "Integrate with Alpaca and Interactive Brokers to execute recommended rebalancing trade tickets with 1 click."),
        ("• Tax-Loss Harvesting & Capital Gains", "Optimize rebalancing frequency to minimize short-term capital gains taxes while keeping portfolio near tangency."),
        ("• Factor-Tilt Constraints (Fama-French)", "Allow investors to constrain portfolios to Value, Momentum, Quality, or Low-Beta multi-factor regimes."),
        ("• Real-Time Intraday Alerts", "Send Discord/Telegram/Email notifications when portfolio weights drift beyond acceptable risk thresholds."),
        ("• ESG & Sustainable Investing Filters", "Integrate ESG risk scores to filter out high-carbon or controversial assets.")
    ]
    for i, (f_head, f_sub) in enumerate(future_items):
        p_h = tf_f.add_paragraph() if i > 0 else tf_f.paragraphs[0]
        p_h.text = f_head
        p_h.font.name = "Arial"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = DARK_GREEN
        p_h.space_before = Pt(6) if i > 0 else Pt(0)

        p_s = tf_f.add_paragraph()
        p_s.text = f_sub
        p_s.font.name = "Arial"
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = DARK_SLATE

    # Right: Project Value
    c_val = add_card(s7, Inches(6.8), Inches(2.0), Inches(5.5), Inches(4.7), LIGHT_GREEN, BORDER_GREEN)
    tf_v = c_val.text_frame
    tf_v.margin_left = Inches(0.3)
    tf_v.margin_top = Inches(0.3)

    val_items = [
        ("• Democratizes Hedge-Fund Math", "Brings \$24k/year institutional tools (Ledoit-Wolf shrinkage, Charnes-Cooper QP, Cornish-Fisher VaR) to everyday retail investors."),
        ("• Eliminates the 'Illusion of Diversification'", "Exposes false diversification before catastrophic drawdowns happen, protecting user savings."),
        ("• Makes Quantitative Analytics Actionable", "Converts dry covariance matrices and Greek letters into exact trade instructions and clear plain-English summaries."),
        ("• Engaging & Educational", "With Health Scores and Roast Mode, learning Modern Portfolio Theory becomes fun, viral, and immediately useful.")
    ]
    for i, (v_head, v_sub) in enumerate(val_items):
        p_h = tf_v.add_paragraph() if i > 0 else tf_v.paragraphs[0]
        p_h.text = v_head
        p_h.font.name = "Arial"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = DARK_GREEN
        p_h.space_before = Pt(8) if i > 0 else Pt(0)

        p_s = tf_v.add_paragraph()
        p_s.text = v_sub
        p_s.font.name = "Arial"
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 8: Conclusion
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8)

    c_concl = add_card(s8, Inches(1.5), Inches(1.2), Inches(10.3), Inches(5.1), WHITE, BORDER_GREEN)
    tf_c = c_concl.text_frame
    tf_c.margin_left = Inches(0.6)
    tf_c.margin_right = Inches(0.6)
    tf_c.margin_top = Inches(0.6)

    p_ct = tf_c.paragraphs[0]
    p_ct.text = "CONCLUSION"
    p_ct.font.name = "Arial Black"
    p_ct.font.size = Pt(32)
    p_ct.font.bold = True
    p_ct.font.color.rgb = DARK_GREEN
    p_ct.space_after = Pt(14)

    concl_paras = [
        "CapitalX bridges the gap between complex quantitative portfolio theory and real-world investor decisions.",
        "By combining automated market data ingestion, robust matrix shrinkage, and convex quadratic programming with an intuitive AI diagnostic layer, CapitalX turns dry statistics into a visual, data-driven decision-support tool.",
        "From identifying hidden correlations to running historical crisis stress tests and generating broker-ready trade tickets, CapitalX empowers investors to build truly resilient, mathematically optimal portfolios with confidence.",
        "Live Terminal Demo: http://localhost:3000 | Backend API Engine: http://localhost:8000"
    ]
    for cp_text in concl_paras:
        p = tf_c.add_paragraph()
        p.text = cp_text
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_SLATE
        p.space_after = Pt(10)

    prs.save("CapitalX_Hackathon_Presentation.pptx")
    print("SUCCESS: CapitalX_Hackathon_Presentation.pptx created successfully!")

if __name__ == "__main__":
    create_deck()
