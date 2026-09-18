import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def build_natural_submission_doc():
    doc = docx.Document()

    # Standard 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Clean formatting
    NAVY = RGBColor(30, 58, 138)       # #1e3a8a
    DARK = RGBColor(31, 41, 55)        # #1f2937
    GRAY = RGBColor(75, 85, 99)        # #4b5563
    ACCENT = RGBColor(37, 99, 235)     # #2563eb

    def add_section_header(title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = NAVY
        
        # Add subtle underline rule
        pPr = h._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '8')
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), 'CBD5E1')
        pBdr.append(bottom)
        pPr.append(pBdr)

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            br = p.add_run(bold_prefix)
            br.font.name = "Arial"
            br.font.size = Pt(10)
            br.font.bold = True
            br.font.color.rgb = DARK
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK
        return p

    def add_bullet(bold_label, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        r_bold = p.add_run(bold_label + ": ")
        r_bold.font.name = "Arial"
        r_bold.font.size = Pt(10)
        r_bold.font.bold = True
        r_bold.font.color.rgb = DARK
        r_text = p.add_run(desc)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = DARK

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("CapitalX: Portfolio Optimization & Risk Intelligence Platform")
    t_run.font.name = "Arial"
    t_run.font.size = Pt(18)
    t_run.font.bold = True
    t_run.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    s_run = sub_p.add_run("Hackathon Idea Submission | Track: Fintech & AI")
    s_run.font.name = "Arial"
    s_run.font.size = Pt(10.5)
    s_run.font.italic = True
    s_run.font.color.rgb = ACCENT

    # 1. Problem Statement
    add_section_header("1. Problem Statement")
    add_body(
        "When we looked at how most individual investors manage their money, we noticed a major recurring problem. "
        "People often buy five or six popular stocks—like Apple, Microsoft, Nvidia, Google, and Amazon—and assume their portfolio is diversified just because they hold different company names. "
        "In reality, these assets share high correlation and move in lockstep. When market headwinds hit the technology sector, the entire portfolio experiences severe drawdowns because the investor only had a single underlying bet."
    )
    add_body(
        "At the same time, existing portfolio analysis tools are not built for normal people. Legacy websites like Portfolio Visualizer were designed years ago for quantitative actuaries. "
        "They present confusing forms and display raw covariance matrices, kurtosis numbers, and statistical jargon without answering the practical questions that matter: "
        "'How risky is my portfolio, what percentage of each stock should I actually own, and what exact trades should I make today?' "
        "We built CapitalX to fix this disconnect by running institutional quantitative optimization under the hood while delivering clear, understandable takeaways."
    )

    # 2. Proposed Solution
    add_section_header("2. Proposed Solution")
    add_body(
        "CapitalX is an interactive web platform where users can test, optimize, and diagnose their investment portfolios. "
        "The platform operates on two interconnected layers: a core quantitative optimization engine and an intelligent diagnostic layer."
    )
    add_body("Core Quantitative Engine (The Foundation):", bold_prefix=None)
    add_bullet("Market Data Ingestion", "The user inputs any set of stock tickers, current portfolio weights, and a historical lookback period (1 to 5 years). The engine downloads adjusted daily closing prices from Yahoo Finance and calculates continuous daily log returns.")
    add_bullet("Statistical Analysis", "Calculates individual annualized expected returns, asset volatilities, and the full pairwise Pearson correlation matrix to reveal which assets are secretly moving together.")
    add_bullet("Covariance Matrix Estimation", "Uses Ledoit-Wolf shrinkage to condition the covariance matrix against sample noise and clips eigenvalues to ensure the matrix remains positive semi-definite.")
    add_bullet("Mean-Variance Portfolio Optimization", "Uses numerical quadratic programming (SciPy SLSQP and CVXPY) to solve for two key benchmark portfolios:\n"
               "  1. Global Minimum Volatility (GMV) Portfolio: Allocates capital to achieve the lowest possible risk.\n"
               "  2. Maximum Sharpe Ratio (Tangency) Portfolio: Finds the exact asset weights that maximize expected return per unit of volatility relative to the risk-free rate.")
    add_bullet("Markowitz Efficient Frontier", "Generates and plots the complete Efficient Frontier curve by sweeping across target returns, showing the user exactly where their current portfolio sits compared to optimal portfolios.")
    add_bullet("Rebalancing & Trade Recommendations", "Compares current weights against optimal tangency weights and produces an asset allocation table with exact percentage changes and recommended actions (BUY, SELL, or HOLD).")

    add_body("Intelligence & Decision-Support Layer (The Practical Edge):", bold_prefix=None)
    add_bullet("Portfolio Health Score (0-100)", "Combines diversification, Sharpe efficiency, volatility discipline, and tail-risk into a single intuitive health grade (A+ to F).")
    add_bullet("Spectral Effective Bets (N_eff)", "Uses Shannon entropy on correlation eigenvalues to show true diversification. If an investor holds 5 tech stocks, it reveals they mathematically only have ~1.3 independent bets.")
    add_bullet("Plain-English AI Commentary", "Translates mathematical results into a readable summary explaining risk blindspots and rebalancing logic, with an optional 'Roast My Portfolio' mode that points out common retail investing mistakes.")
    add_bullet("1-Click Crisis Simulator", "Stress-tests the portfolio through historical shocks like the 2020 COVID crash, 2022 Fed rate hikes, and the 2008 financial crisis, showing simulated drawdowns and recovery time.")
    add_bullet("Interactive 'What-If' Sandbox", "Allows users to test adding non-correlated assets like Gold, US Treasuries, or Cash with one click to see the impact on their frontier before committing real money.")
    add_bullet("1-Click Institutional Tear Sheet", "Allows users to export a clean, printable 1-page investment summary PDF.")

    # 3. Technology & Quantitative Architecture
    add_section_header("3. Technology & Architecture")
    add_bullet("Backend Framework", "Python 3.11 with FastAPI. Handles asynchronous API requests with sub-second response times.")
    add_bullet("Quantitative Libraries", "NumPy and Pandas for fast return series manipulation. SciPy and CVXPY for quadratic optimization solvers (SLSQP, OSQP, Clarabel) with multi-tier fallbacks to prevent optimization failures.")
    add_bullet("Data Pipeline", "yfinance integration with thread-safe in-memory caching to minimize redundant external API calls.")
    add_bullet("Frontend Stack", "Vanilla modern JavaScript (ES6+ modules), HTML5, and CSS3. We avoided heavy frontend frameworks like React to keep bundle size near zero and ensure fast load times.")
    add_bullet("Visualizations", "Chart.js for interactive Efficient Frontier scatter plots, correlation heatmaps, asset volatility comparisons, and rolling risk graphs.")

    # 4. Innovation & What Makes CapitalX Different
    add_section_header("4. Innovation & Uniqueness")
    add_bullet("Actionable Trade Instructions", "Instead of stopping at charts and formulas, CapitalX tells the user the exact percentage adjustment needed for each stock to reach optimal risk-adjusted returns.")
    add_bullet("True Diversification vs Nominal Diversification", "Most apps just count how many stocks you own. We calculate the spectral Effective Number of Bets to expose false diversification.")
    add_bullet("Accessible Risk Communication", "By translating Greek statistics into a clear Health Score and plain-English memos, anyone from a student to an experienced retail trader can immediately understand their portfolio risk.")
    add_bullet("Interactive Experimentation", "The What-If sandbox lets users explore realistic hedging strategies in seconds without needing spreadsheet modeling skills.")

    # 5. Impact & Future Roadmap
    add_section_header("5. Impact & Future Roadmap")
    add_bullet("Democratizing Institutional Finance", "Gives everyday investors access to portfolio optimization techniques that were previously only available in costly enterprise software like Bloomberg or Morningstar Direct.")
    add_bullet("Preventing Avoidable Capital Losses", "Helps investors spot excessive correlation and risk concentration before market downturns happen.")
    add_bullet("Future Enhancements", "We plan to integrate direct broker APIs (such as Alpaca and Interactive Brokers) so users can execute rebalancing trades directly from the dashboard, as well as adding ESG and factor-tilt constraints.")

    target_file = "CapitalX_Hackathon_Idea_Submission.docx"
    try:
        doc.save(target_file)
        print(f"SUCCESS: {target_file} generated successfully!")
    except PermissionError:
        fallback_file = "CapitalX_Hackathon_Submission.docx"
        doc.save(fallback_file)
        print(f"Original file was open in Word. Saved to fallback: {fallback_file}")

if __name__ == "__main__":
    build_natural_submission_doc()
